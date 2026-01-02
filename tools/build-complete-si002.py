#!/usr/bin/env python3
"""Build complete SI002 HTML template from all content"""

from docx import Document
from pathlib import Path
import re

doc_path = Path(__file__).parent.parent / "SI002 - SII Document Request - Triad - V1.0.docx"
output_path = Path(__file__).parent.parent / "formatter examples" / "SI002" / "SI002-formatted.html"

doc = Document(str(doc_path))
paras = [p for p in doc.paragraphs if p.text.strip()]

# Find main content start and end
main_idx = next(i for i, p in enumerate(paras) if 'You may qualify' in p.text)
closing_idx = next(i for i, p in enumerate(paras) if 'Sincerely' in p.text or 'You may obtain' in p.text)

# Extract all content paragraphs
content_paras = []
for i, p in enumerate(paras[main_idx:closing_idx+5]):
    text = p.text.strip()
    if text and len(text) > 2:
        # Skip pure metadata
        if not (text.startswith('[') and ']' in text and len(text) < 50):
            content_paras.append(text)

# Build HTML
html_lines = [
    '<div>{Insert(H003 TagHeader)}</div>',
    '<br>',
    '<div>{[L001]}</div>',
    '<div>{[mailingAddress]}</div>',
    '<br><br><br><br><br>',
    '<table width="100%"><tbody><tr>',
    '  <td width="20%" valign="top">Re: Loan Number:</td>',
    '  <td>{[M594]}</td>',
    '</tr><tr>',
    '  <td width="20%" valign="top">RE:</td>',
    '  <td>{Compress({[M567]}|{[M583]}|{[M568]})}</td>',
    '</tr></tbody></table>',
    '<br>',
    '<div>Dear {[Salutation]},</div>',
    '<br>'
]

# Process content paragraphs
i = 0
while i < len(content_paras):
    text = content_paras[i]
    
    # Check for IF statements
    if text.startswith('IF M960'):
        # Extract state condition
        match = re.search(r'IF M960.*?= (.+?) THEN:', text)
        if match:
            states_str = match.group(1).strip()
            # Convert to proper If() syntax
            if ',' in states_str:
                # Multiple states - use IN
                states = [s.strip() for s in states_str.split(',')]
                states_str = "', '".join(states)
                html_lines.append(f"{{If('{{[M960]}}' IN ('{states_str}'))}}")
            else:
                # Single state
                state = states_str.strip()
                html_lines.append(f"{{If('{{[M960]}}' = '{state}')}}")
        i += 1
        continue
    
    # Regular content
    # Replace variable references
    text = re.sub(r'\[M567\]', '{[M567]}', text)
    text = re.sub(r'\[M583\]', '{[M583]}', text)
    text = re.sub(r'\[M568\]', '{[M568]}', text)
    text = re.sub(r'\[M594\]', '{[M594]}', text)
    text = re.sub(r'\[M960\]', '{[M960]}', text)
    
    # Replace company placeholders
    text = re.sub(r'<CSFax>', '{[plsMatrix.CSFax]}', text)
    text = re.sub(r'<CSEmail>', '{[plsMatrix.CSEmail]}', text)
    text = re.sub(r'<CSPhoneNumber>', '{[plsMatrix.CSPhoneNumber]}', text)
    text = re.sub(r'<CompanyLongName>', '{[plsMatrix.CompanyLongName]}', text)
    text = re.sub(r'<CompanyReturnAddr1>', '{[plsMatrix.CompanyReturnAddr1]}', text)
    text = re.sub(r'<CompanyReturnAddr2>', '{[plsMatrix.CompanyReturnAddr2]}', text)
    text = re.sub(r'<CompanyReturnAddr3>', '{[plsMatrix.CompanyReturnAddr3]}', text)
    
    # Replace "Triad Financial Services, Inc." with plsMatrix
    text = re.sub(r'Triad Financial Services, Inc\.', '{[plsMatrix.CompanyLongName]}', text)
    
    # Add as div
    html_lines.append(f'<div>{text}</div>')
    html_lines.append('<br>')
    
    # Check if next is End If (when we hit a new IF or closing)
    if i + 1 < len(content_paras):
        next_text = content_paras[i + 1]
        if next_text.startswith('IF M960') or 'Please review' in next_text:
            # Need to close current conditional block
            html_lines.append('{End If}')
            html_lines.append('<br>')
    
    i += 1

# Add closing
html_lines.extend([
    '<br>',
    '<div>Please review the circumstances listed above, and provide us with the appropriate documentation for the option that best describes your situation via fax to {[plsMatrix.CSFax]}, email to {[plsMatrix.CSEmail]}, or by mail to:</div>',
    '<br>',
    '<div>{[plsMatrix.CompanyLongName]}</div>',
    '<div>{[plsMatrix.CompanyReturnAddr1]}</div>',
    '<div>{[plsMatrix.CompanyReturnAddr2]}</div>',
    '<div>{[plsMatrix.CompanyReturnAddr3]}</div>',
    '<br>',
    '<div>You may obtain a more individualized description of required documents by providing additional information. If you have any questions or concerns regarding this request, please contact our Customer Service Department at {[plsMatrix.CSPhoneNumber]}.</div>',
    '<br>',
    '<div>Sincerely,</div>',
    '<br>',
    '<div>Customer Service Department</div>',
    '<div>{[plsMatrix.CompanyLongName]}</div>'
])

# Write output
output_path.parent.mkdir(parents=True, exist_ok=True)
with open(output_path, 'w', encoding='utf-8') as f:
    f.write('\n'.join(html_lines))

print(f"Generated template with {len(html_lines)} lines")
print(f"Saved to: {output_path}")

